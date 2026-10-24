#!/usr/bin/env python3
"""
scripts/batch_clean_sources.py

High-performance automated batch file cleaning pipeline for all vendor deed source files
registered in `client/src/data/sourcesData.js`.

Pipeline Architecture:
1. Multi-threaded batch engine powered by `python-calamine` (Rust engine), with `openpyxl`
   and `xlrd` fallbacks for edge-case workbooks.
2. Sources location:
   - 1,118 unique files pre-cached locally at `C:\\Users\\Rishi D\\Downloads\\NEW_DATA_2026_EXTRACTED`
   - Missing / on-demand files (e.g. `TownSquare.xlsx`) downloaded automatically from Supabase Storage:
     `https://qgxgtavkovqklijfpnfl.supabase.co/storage/v1/object/public/dubai_real_estate/sources/<filename>`
3. Dynamic header detection: scans up to the first 10 rows for column keywords, penalizing data rows
   to skip title banners and offset headers (e.g. `MUDON p1.xls` row 2).
4. Multi-sheet support:
   - Pairs and joins relational `OWNER` + `PROPERTY` sheets on `P-NUMBER` (e.g. `Wadi Al Safa 5.xlsx`).
   - Parses all data sheets across both `.xlsx` and legacy binary `.xls` workbooks.
   - Skips empty worksheets.
5. Normalization:
   - Phone numbers -> E.164 (+971...) standard: numeral translation, strip non-digits, fix +97105 to +9715,
     add +971 to 9-digit local numbers, multi-phone cell splitting (extract primary mobile).
   - Unit & DEWA Parser -> strips DEWA pipe compound strings, cadastral plot prefixes, isolates unit number.
     Heals or rejects cell-shifted dates or phone numbers in unit columns.
   - Owner Name Sanitizer -> cleans names, reverses "Last, First", title-cases, handles corporate acronyms
     (LLC, PJSC, DMCC), falls back to "Property Owner".
   - Disambiguates project, building, and community references.
6. Target schema:
   Exact 9 uniform RFC 4180 CSV columns:
   full_name,phone_number_normalized,unit_number,property_reference,property_location,property_type,bedrooms,size_sqft,status
7. Output destination:
   `cleaned_sources_output/<filename_base>_cleaned.csv`
8. Manifest:
   `cleaning_manifest.json` recording for each unique file: status, original rows, cleaned rows, duration, output path.
"""

import os
import sys
import re
import csv
import time
import json
import requests
from concurrent.futures import ThreadPoolExecutor, as_completed

# Force UTF-8 stdout on Windows
if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

import python_calamine
import openpyxl
import xlrd

# Paths
BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
SOURCES_DATA_JS = os.path.join(BASE_DIR, "client", "src", "data", "sourcesData.js")
LOCAL_DIR = r"C:\Users\Rishi D\Downloads\NEW_DATA_2026_EXTRACTED"
OUTPUT_DIR = os.path.join(BASE_DIR, "cleaned_sources_output")
CACHE_DIR = os.path.join(BASE_DIR, "download_cache")
MANIFEST_PATH = os.path.join(BASE_DIR, "cleaning_manifest.json")
SUPABASE_SOURCES_URL = "https://qgxgtavkovqklijfpnfl.supabase.co/storage/v1/object/public/dubai_real_estate/sources"

os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(CACHE_DIR, exist_ok=True)

# Domain Knowledge Sets
DUMMY_STRINGS = {
    'null', 'n/a', 'na', 'none', '-', '--', '0', 'undefined', 'nil', 'unknown',
    '.', 'n.a', 'not available', 'table 1', 'unnamed', 'owner', 'tenant', 'nan',
    'residential', 'commercial', 'dubai', 'uae', 'u.a.e', 'not assigned', 'p.o. box'
}

ACRONYMS = {'LLC', 'PJSC', 'P.J.S.C', 'FZ-LLC', 'FZCO', 'DMCC', 'FZE', 'LTD'}
ARABIC_INDIC_MAP = str.maketrans('٠١٢٣٤٥٦٧٨٩', '0123456789')

COMMUNITY_MAP = {
    'palm': 'Palm Jumeirah',
    'palm jumeirah': 'Palm Jumeirah',
    'jbr': 'Jumeirah Beach Residence (JBR)',
    'jumeirah beach residence': 'Jumeirah Beach Residence (JBR)',
    'marina': 'Dubai Marina',
    'dubai marina': 'Dubai Marina',
    'marsa dubai': 'Dubai Marina',
    'downtown': 'Downtown Dubai',
    'burj khalifa': 'Downtown Dubai',
    'business bay': 'Business Bay',
    'jvc': 'Jumeirah Village Circle (JVC)',
    'jumeirah village circle': 'Jumeirah Village Circle (JVC)',
    'jvt': 'Jumeirah Village Triangle (JVT)',
    'jumeirah village triangle': 'Jumeirah Village Triangle (JVT)',
    'jlt': 'Jumeirah Lake Towers (JLT)',
    'jumeirah lake towers': 'Jumeirah Lake Towers (JLT)',
    'jumeirah islands': 'Jumeirah Islands',
    'jumeirah park': 'Jumeirah Park',
    'jumeirah': 'Jumeirah',
    'difc': 'DIFC',
    'central park': 'DIFC',
    'mudon': 'Mudon',
    'al furjan': 'Al Furjan',
    'furjan': 'Al Furjan',
    'damac hills': 'Damac Hills',
    'damac lagoons': 'Damac Lagoons',
    'damac lagoon': 'Damac Lagoons',
    'akoya': 'Damac Hills 2 (Akoya)',
    'dubai hills': 'Dubai Hills Estate',
    'dubai hills estate': 'Dubai Hills Estate',
    'creek': 'Dubai Creek Harbour',
    'dubai creek harbour': 'Dubai Creek Harbour',
    'al barsha': 'Al Barsha',
    'al barshaa': 'Al Barsha',
    'barsha': 'Al Barsha',
    'arabian ranches': 'Arabian Ranches',
    'ranches': 'Arabian Ranches',
    'townsquare': 'Town Square',
    'town square': 'Town Square',
    'meydan': 'Meydan (MBR City)',
    'mbr': 'Mohammed Bin Rashid City (MBR City)',
    'mohammed bin rashid': 'Mohammed Bin Rashid City (MBR City)',
    'sports city': 'Dubai Sports City',
    'motor city': 'Motor City',
    'studio city': 'Dubai Studio City',
    'silicon': 'Dubai Silicon Oasis',
    'dso': 'Dubai Silicon Oasis',
    'international city': 'International City',
    'warsan': 'International City',
    'al satwa': 'Al Satwa',
    'satwa': 'Al Satwa',
    'al quoz': 'Al Quoz',
    'al safa': 'Al Safa',
    'wadi al safa': 'Wadi Al Safa',
    'nad al shiba': 'Nad Al Sheba',
    'nad al sheba': 'Nad Al Sheba',
    'mirdif': 'Mirdif',
    'al qusais': 'Al Qusais',
    'deira': 'Deira',
    'bur dubai': 'Bur Dubai',
    'tecom': 'Barsha Heights (Tecom)',
    'barsha heights': 'Barsha Heights (Tecom)',
    'the greens': 'The Greens',
    'the views': 'The Views',
    'the lakes': 'The Lakes',
    'meadows': 'The Meadows',
    'springs': 'The Springs',
    'emirates hills': 'Emirates Hills',
    'tilal al ghaf': 'Tilal Al Ghaf',
    'sobha hartland': 'Sobha Hartland',
    'hartland': 'Sobha Hartland',
    'al kifaf': 'Al Kifaf',
    'zaabeel': 'Za\'abeel',
    'city walk': 'City Walk',
    'bluewaters': 'Bluewaters Island',
    'mina rashid': 'Mina Rashid',
    'port de la mer': 'Port De La Mer',
    'la mer': 'La Mer',
    'al habtoor': 'Al Habtoor City',
    'habtoor': 'Al Habtoor City',
    'villanova': 'Villanova',
    'serena': 'Serena',
    'cherrywoods': 'Cherrywoods',
    'the oasis': 'The Oasis',
    'mirage': 'The Oasis',
    'athlon': 'Athlon',
}

KEYWORD_TERMS = {
    'name', 'owner', 'client', 'customer', 'buyer', 'seller',
    'phone', 'mobile', 'contact', 'cell', 'tel', 'whatsapp',
    'unit', 'flat', 'villa', 'apartment', 'suite', 'premise', 'plot',
    'project', 'building', 'tower', 'location', 'area', 'community',
    'type', 'usage', 'bedrooms', 'bed', 'beds', 'rooms', 'size', 'sqft',
    'p-number', 'p number', 'p_number', 'registration'
}

NEG_PHONE = {'land', 'plot', 'flat', 'unit', 'villa', 'building', 'door', 'premise', 'procedure', 'id', 'license', 'licence', 'emirates', 'passport', 'box', 'post', 'pnumber', 'registration', 'amount', 'price', 'size', 'date', 'fee', 'value'}
NEG_NAME = {'building', 'project', 'tower', 'community', 'company', 'agency', 'broker', 'bank', 'zone', 'procedure', 'id', 'code', 'date', 'amount', 'size', 'price', 'fee', 'unit', 'flat', 'plot', 'phone', 'mobile', 'tel', 'whatsapp', 'email'}
NEG_UNIT = {'type', 'size', 'price', 'amount', 'date', 'status', 'bed', 'beds', 'phone', 'mobile', 'email', 'name', 'buyer', 'seller', 'client', 'customer'}
NEG_LOC = {'size', 'totalarea', 'actualarea', 'commonarea', 'balconyarea', 'amount', 'price', 'phone', 'date', 'name'}

def has_neg(h, neg_set):
    if h.startswith('unnamed'):
        return False
    return any(neg in h for neg in neg_set)


# --- Normalization Utilities ---

def extract_phone_e164(val):
    if val is None:
        return ""
    if isinstance(val, (int, float)):
        s = f"{val:.0f}" if isinstance(val, float) and val.is_integer() else str(val).strip()
    else:
        s = str(val).strip()
    
    if not s or s.lower() in ('nan', 'none', 'null', '0', '0.0', 'undefined', 'n/a', '-'):
        return ""
    
    s = s.translate(ARABIC_INDIC_MAP)
    if '@' in s or (s.count('-') >= 2 and '00:00' in s):
        return ""

    s = re.sub(r'\b\d+\.?\d*E[+-]?\d+\b', ' ', s, flags=re.I)
    
    mob_match = re.search(r'(?:mobile|mob|cell)\s*\d*[:\s-]*([^,;\n\r|&]+)', s, re.I)
    if mob_match:
        cand = mob_match.group(1).strip()
        cleaned_cand = find_best_phone_in_text(cand)
        if cleaned_cand and re.match(r'^\+[1-9]\d{6,14}$', cleaned_cand):
            return cleaned_cand

    res = find_best_phone_in_text(s)
    if res and re.match(r'^\+[1-9]\d{6,14}$', res):
        return res
    return ""

def find_best_phone_in_text(text):
    if not text:
        return ""
    cleaned_text = re.sub(r'^(?:\+?00971|\+?971|00971)[\s|/\\-]+0?', '971', text)
    cleaned_text = re.sub(r'[\s|/\\-]+(?:\+?00971|\+?971|00971)[\s|/\\-]+0?', ' 971', cleaned_text)

    # 1. +971 or 971 followed by 5 and 8 digits
    uae_intl = re.findall(r'(?:\+?00971|\+?971|00971)[\s\-.]*(?:0)?(5\d[\s\-.]*\d{3}[\s\-.]*\d{4})', cleaned_text)
    if uae_intl:
        digits = re.sub(r'\D', '', uae_intl[0])
        if len(digits) == 9:
            return f"+971{digits}"

    # 2. Local 05x (10 digits)
    uae_local_10 = re.findall(r'(?<!\d)0(5\d[\s\-.]*\d{3}[\s\-.]*\d{4})(?!\d)', cleaned_text)
    if uae_local_10:
        digits = re.sub(r'\D', '', uae_local_10[0])
        if len(digits) == 9:
            return f"+971{digits}"

    # 3. Local 5x (9 digits)
    uae_local_9 = re.findall(r'(?<!\d)(5\d[\s\-.]*\d{3}[\s\-.]*\d{4})(?!\d)', cleaned_text)
    if uae_local_9:
        digits = re.sub(r'\D', '', uae_local_9[0])
        if len(digits) == 9:
            return f"+971{digits}"

    # 4. UAE Landlines (02, 03, 04, 06, 07, 09)
    uae_landline = re.findall(r'(?:\+?00971|\+?971|00971)[\s\-.]*(?:0)?([234679]\d{7})', cleaned_text)
    if uae_landline:
        digits = re.sub(r'\D', '', uae_landline[0])
        return f"+971{digits}"
    uae_local_land = re.findall(r'(?<!\d)0([234679]\d{7})(?!\d)', cleaned_text)
    if uae_local_land:
        digits = re.sub(r'\D', '', uae_local_land[0])
        return f"+971{digits}"

    # 5. General International (+ or 00)
    intl = re.findall(r'(?:\+|00)(\d{7,15})', re.sub(r'[\s\-.]', '', cleaned_text))
    if intl:
        d = intl[0]
        if not re.match(r'^0+$', d):
            if d.startswith('9710') and len(d) in (12, 13, 14):
                d = '971' + d[4:]
            if d.startswith('05') and len(d) == 10:
                return f"+971{d[1:]}"
            if re.match(r'^0[234679]\d{7}$', d):
                return f"+971{d[1:]}"
            if d.startswith('5') and len(d) == 9:
                return f"+971{d}"
            if re.match(r'^[234679]\d{7}$', d):
                return f"+971{d}"
            if d.startswith('0'):
                d = d.lstrip('0')
            if 7 <= len(d) <= 15 and not d.startswith('0') and not re.match(r'^0+$', d):
                return f"+{d}"

    chunks = re.split(r'[,;\n\r/|&]+', cleaned_text)
    for chunk in chunks:
        d = re.sub(r'\D', '', chunk)
        if d.startswith('9710') and len(d) in (12, 13, 14):
            d = '971' + d[4:]
        if d.startswith('05') and len(d) == 10:
            return f"+971{d[1:]}"
        if re.match(r'^0[234679]\d{7}$', d):
            return f"+971{d[1:]}"
        if d.startswith('5') and len(d) == 9:
            return f"+971{d}"
        if re.match(r'^[234679]\d{7}$', d):
            return f"+971{d}"
        if d.startswith('0'):
            d = d.lstrip('0')
        if 7 <= len(d) <= 15:
            if d.startswith('971'):
                return f"+{d}"
            if len(d) == 8 and d[0] in '234679':
                return f"+971{d}"
            if len(d) == 9 and d[0] == '5':
                return f"+971{d}"
            if not d.startswith('0') and not re.match(r'^0+$', d):
                return f"+{d}"

    return ""

def clean_owner_name(raw):
    if raw is None:
        return "Property Owner"
    s = str(raw).strip()
    if not s or s.lower() in DUMMY_STRINGS:
        return "Property Owner"
        
    if '@' in s and '.' in s:
        return "Property Owner"
    if re.match(r'^\+?[\d\s\-_/()]{7,}$', s):
        return "Property Owner"
    if re.match(r'^\d{4}[-/]\d{1,2}[-/]\d{1,2}', s):
        return "Property Owner"
        
    s = re.sub(r'^(?:NR\/NHO|REF|ID|CODE)[\s,:\-_]+', '', s, flags=re.I)
    s = re.sub(r'-(?:Kempinski|Tower|Building|Villa|Flat|Unit|TH)[\s\w]*$', '', s, flags=re.I)
    
    # Corporate acronym normalization
    s = re.sub(r'\(?\s*L[\s.]*L[\s.]*C\.?\s*\)?', ' LLC ', s, flags=re.I)
    s = re.sub(r'\(?\s*P[\s.]*J[\s.]*S[\s.]*C\.?\s*\)?', ' P.J.S.C ', s, flags=re.I)
    
    s = re.sub(r'[\\/\[\]{}()_]', ' ', s)
    s = re.sub(r'\s+-\s+', ' ', s)
    s = re.sub(r'^[\s,.;:-]+|[\s,.;:-]+$', '', s)
    
    # Handle inverted / comma separated names
    if ',' in s:
        parts = [p.strip() for p in s.split(',') if p.strip()]
        upper = s.upper()
        has_corp = any(a in upper for a in ACRONYMS) or re.search(r'\b(?:LLC|LIMITED|LTD|PROPERTIES|DEVELOPMENT|HOLDING|INVESTMENT)\b', upper)
        if len(parts) == 2 and not has_corp and '&' not in upper and ' AND ' not in upper:
            p0, p1 = parts[0], parts[1]
            p0_low, p1_low = p0.lower(), p1.lower()
            if re.search(r'(?:ova|eva|ina|aya|ov|ev|in|skiy|sky)$', p1_low):
                s = f"{p0} {p1}"
            else:
                s = f"{p1} {p0}"
        else:
            s = ' '.join(parts)
            
    s = re.sub(r';', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip()
    if not s or s.lower() in DUMMY_STRINGS:
        return "Property Owner"
        
    words = s.split()
    formatted = []
    for w in words:
        clean_w = re.sub(r'[^a-zA-Z.]', '', w)
        up_w = clean_w.upper()
        if up_w in ACRONYMS:
            formatted.append('P.J.S.C' if up_w == 'PJSC' else up_w)
            continue
        low_w = w.lower()
        if low_w in ('al', 'el'):
            formatted.append('Al')
        elif low_w in ('bin', 'ibn'):
            formatted.append('Bin')
        elif low_w == 'bint':
            formatted.append('Bint')
        elif low_w == 'abu':
            formatted.append('Abu')
        else:
            formatted.append(w.capitalize())
            
    res = ' '.join(formatted).strip()
    return res or "Property Owner"

def clean_unit_number(raw_unit, building_name=''):
    if raw_unit is None:
        return ""
    if isinstance(raw_unit, float) and raw_unit.is_integer():
        u = str(int(raw_unit))
    else:
        u = str(raw_unit).strip()
        if u.endswith('.0') and u[:-2].isdigit():
            u = u[:-2]
            
    if not u or u.lower() in DUMMY_STRINGS:
        return ""
        
    # Reject cell shift: phone number in unit
    digits = re.sub(r'\D', '', u)
    if len(digits) >= 9 and (u.startswith('05') or u.startswith('+971') or u.startswith('971')):
        return ""
        
    # Reject cell shift: date in unit
    if re.match(r'^\d{4}[-/]\d{1,2}[-/]\d{1,2}', u) or re.match(r'^\d{1,2}[-/]\d{1,2}[-/]\d{4}', u):
        return ""
        
    # DEWA pipe string
    if '|' in u:
        parts = [p.strip() for p in u.split('|') if p.strip()]
        apt_part = None
        for p in reversed(parts):
            m = re.search(r'(?:Apartment|Flat|Unit|Villa|TH|Suite|Office|Shop)\s*[:#-]?\s*([A-Za-z0-9\-]+)', p, re.I)
            if m:
                apt_part = m.group(1).strip()
                break
            m2 = re.search(r'\b(?:OFFICE|TOWER\s+[A-Za-z]|BLDG\s+[0-9\-]+)?\s*(\d{2,5}[A-Za-z]?)\b', p, re.I)
            if m2:
                apt_part = m2.group(1).strip()
                break
        if apt_part:
            u = apt_part
        elif parts:
            u = parts[-1]

    u = re.sub(r'^(?:Apartment|Flat|Unit|Villa|Apt|Suite|Door|Premise|Premises|Townhouse|TH)\s*[:#-]?\s*', '', u, flags=re.I)
    
    if building_name and len(building_name) > 3:
        b_clean = building_name.strip()
        if u.upper().startswith(b_clean.upper()):
            u = u[len(b_clean):].lstrip(' -_/:#')

    u = re.sub(r'-(?:0|00)$', '', u)
    return u.strip()

def clean_bedrooms(raw_beds):
    if raw_beds is None:
        return ""
    s = str(raw_beds).strip()
    if not s or s.lower() in DUMMY_STRINGS:
        return ""
    if 'studio' in s.lower():
        return "Studio"
    m = re.search(r'\b(\d{1,2})\b', s)
    if m:
        return m.group(1)
    return s

def clean_size_sqft(raw_size, is_sqm=False):
    if raw_size is None:
        return ""
    s = str(raw_size).strip()
    if not s or s.lower() in DUMMY_STRINGS:
        return ""
    s_clean = re.sub(r'[^\d.]', '', s)
    if not s_clean:
        return ""
    try:
        val = float(s_clean)
        if val <= 0:
            return ""
        if is_sqm or val < 40: # likely sqm
            val = val * 10.7639
        return f"{val:.1f}" if val % 1 != 0 else f"{int(val)}"
    except Exception:
        return ""

def detect_header_row(rows, max_scan=10):
    if not rows:
        return -1, []
    
    best_row_idx = 0
    best_score = -1
    
    for idx in range(min(len(rows), max_scan)):
        row = rows[idx]
        if not row:
            continue
        non_empty = [c for c in row if c is not None and str(c).strip() != '']
        if not non_empty:
            continue
            
        matches = 0
        all_strings = 0
        all_numeric = 0
        has_data_signals = 0
        
        for c in non_empty:
            s = str(c).strip().lower()
            if '@' in s or re.search(r'\+?971|\b05\d{8}\b', s) or '{' in s or (s.count('-') >= 2 and '00:00' in s):
                has_data_signals += 1
            if 'table 1' in s or 'unnamed' in s:
                matches += 1
                
            s_clean = re.sub(r'[^a-z0-9]', ' ', s)
            words = set(s_clean.split())
            if any(k in words or any(s_clean.startswith(k) or s_clean.endswith(k) for k in KEYWORD_TERMS) for k in KEYWORD_TERMS):
                matches += 1
            if re.search(r'[a-zA-Z]', s):
                all_strings += 1
            elif re.match(r'^\d+(\.\d+)?$', s):
                all_numeric += 1
        
        score = (matches * 15) + (all_strings * 2) - (all_numeric * 5) - (has_data_signals * 50)
        if len(non_empty) <= 2 and len(row) > 3 and matches == 0:
            score -= 50
            
        if score > best_score:
            best_score = score
            best_row_idx = idx

    if best_score <= 0:
        best_row_idx = 0

    headers = [str(c).strip() if c is not None and str(c).strip() != '' else f'col_{i}' 
               for i, c in enumerate(rows[best_row_idx])]
    return best_row_idx, headers

def disambiguate_metadata_from_filename(filename):
    clean = re.sub(r'\.(xlsx|xls)$', '', filename, flags=re.I)
    clean = re.sub(r'^\d+[\s._-]+', '', clean)
    clean = re.sub(r'\(\d+\s*contacts?\)', '', clean, flags=re.I)
    clean = re.sub(r'[_\-]+', ' ', clean).strip()
    
    clean_low = clean.lower()
    
    detected_loc = "Dubai"
    for k, v in COMMUNITY_MAP.items():
        if k in clean_low:
            detected_loc = v
            break
            
    detected_type = "Apartment"
    if 'villa' in clean_low:
        detected_type = "Villa"
    elif 'townhouse' in clean_low or 'town house' in clean_low:
        detected_type = "Townhouse"
    elif 'penthouse' in clean_low:
        detected_type = "Penthouse"
    elif 'plot' in clean_low:
        detected_type = "Plot"
    elif 'commercial' in clean_low:
        detected_type = "Commercial"
        
    detected_ref = clean.title() if clean else "Primary Asset"
    return detected_loc, detected_ref, detected_type

def process_single_workbook(filepath, filename):
    t0 = time.time()
    clean_base = re.sub(r'\.(xlsx|xls)$', '', filename, flags=re.I)
    out_csv = os.path.join(OUTPUT_DIR, f"{clean_base}_cleaned.csv")
    
    # Check if empty / lock file
    sz = os.path.getsize(filepath) if os.path.exists(filepath) else 0
    if sz < 500 or filename.startswith('~'):
        with open(out_csv, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow(['full_name', 'phone_number_normalized', 'unit_number', 'property_reference', 'property_location', 'property_type', 'bedrooms', 'size_sqft', 'status'])
        return {
            'filename': filename,
            'clean_base_name': clean_base,
            'status': 'quarantined',
            'original_rows': 0,
            'cleaned_rows': 0,
            'duration': round(time.time() - t0, 3),
            'output_path': out_csv,
            'reason': 'Lock or corrupted Excel file (< 500 bytes)'
        }
        
    # Read sheets
    sheet_data = {}
    try:
        wb = python_calamine.CalamineWorkbook.from_path(filepath)
        for sname in wb.sheet_names:
            sheet = wb.get_sheet_by_name(sname)
            rows = sheet.to_python()
            if rows:
                sheet_data[sname] = rows
    except Exception as cal_err:
        # Fallback to xlrd or openpyxl
        if filepath.lower().endswith('.xls'):
            try:
                xwb = xlrd.open_workbook(filepath)
                for sname in xwb.sheet_names():
                    xsheet = xwb.sheet_by_name(sname)
                    rows = [xsheet.row_values(rx) for rx in range(xsheet.nrows)]
                    if rows:
                        sheet_data[sname] = rows
            except Exception as xlrd_err:
                with open(out_csv, 'w', newline='', encoding='utf-8') as f:
                    writer = csv.writer(f)
                    writer.writerow(['full_name', 'phone_number_normalized', 'unit_number', 'property_reference', 'property_location', 'property_type', 'bedrooms', 'size_sqft', 'status'])
                return {'filename': filename, 'clean_base_name': clean_base, 'status': 'error', 'original_rows': 0, 'cleaned_rows': 0, 'duration': round(time.time() - t0, 3), 'output_path': out_csv, 'reason': f"Cal: {cal_err}, xlrd: {xlrd_err}"}
        else:
            try:
                owb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
                for sname in owb.sheetnames:
                    ws = owb[sname]
                    rows = list(ws.iter_rows(values_only=True))
                    if rows:
                        sheet_data[sname] = rows
                owb.close()
            except Exception as opx_err:
                with open(out_csv, 'w', newline='', encoding='utf-8') as f:
                    writer = csv.writer(f)
                    writer.writerow(['full_name', 'phone_number_normalized', 'unit_number', 'property_reference', 'property_location', 'property_type', 'bedrooms', 'size_sqft', 'status'])
                return {'filename': filename, 'clean_base_name': clean_base, 'status': 'error', 'original_rows': 0, 'cleaned_rows': 0, 'duration': round(time.time() - t0, 3), 'output_path': out_csv, 'reason': f"Cal: {cal_err}, opx: {opx_err}"}

    if not sheet_data:
        with open(out_csv, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow(['full_name', 'phone_number_normalized', 'unit_number', 'property_reference', 'property_location', 'property_type', 'bedrooms', 'size_sqft', 'status'])
        return {
            'filename': filename,
            'clean_base_name': clean_base,
            'status': 'quarantined',
            'original_rows': 0,
            'cleaned_rows': 0,
            'duration': round(time.time() - t0, 3),
            'output_path': out_csv,
            'reason': 'Empty workbook or no sheets with data'
        }

    default_loc, default_ref, default_type = disambiguate_metadata_from_filename(filename)

    # Check for OWNER + PROPERTY paired sheets
    upper_sheets = {k.upper(): k for k in sheet_data.keys()}
    prop_lookup = {}
    if 'OWNER' in upper_sheets and 'PROPERTY' in upper_sheets:
        prop_rows = sheet_data[upper_sheets['PROPERTY']]
        p_hdr_idx, p_hdrs = detect_header_row(prop_rows)
        p_hdr_lower = [re.sub(r'[^a-z0-9]', '', str(h).lower()) for h in p_hdrs]
        p_key_idx = next((i for i, h in enumerate(p_hdr_lower) if 'pnumber' in h or h == 'id'), -1)
        p_flat_idx = next((i for i, h in enumerate(p_hdr_lower) if 'flat' in h or 'unit' in h), -1)
        p_bldg_idx = next((i for i, h in enumerate(p_hdr_lower) if 'building' in h or 'project' in h), -1)
        p_area_idx = next((i for i, h in enumerate(p_hdr_lower) if 'area' in h or 'master' in h), -1)
        p_type_idx = next((i for i, h in enumerate(p_hdr_lower) if 'propertytype' in h or 'usage' in h), -1)
        p_rooms_idx = next((i for i, h in enumerate(p_hdr_lower) if 'rooms' in h or 'bedroom' in h), -1)
        p_size_idx = next((i for i, h in enumerate(p_hdr_lower) if 'actualarea' in h or 'size' in h or 'totalarea' in h), -1)

        if p_key_idx != -1:
            for r in prop_rows[p_hdr_idx + 1:]:
                if len(r) > p_key_idx and r[p_key_idx] is not None:
                    k = str(r[p_key_idx]).strip()
                    if k:
                        prop_lookup[k] = {
                            'unit': str(r[p_flat_idx]).strip() if p_flat_idx != -1 and len(r) > p_flat_idx and r[p_flat_idx] is not None else '',
                            'building': str(r[p_bldg_idx]).strip() if p_bldg_idx != -1 and len(r) > p_bldg_idx and r[p_bldg_idx] is not None else '',
                            'location': str(r[p_area_idx]).strip() if p_area_idx != -1 and len(r) > p_area_idx and r[p_area_idx] is not None else '',
                            'property_type': str(r[p_type_idx]).strip() if p_type_idx != -1 and len(r) > p_type_idx and r[p_type_idx] is not None else '',
                            'bedrooms': str(r[p_rooms_idx]).strip() if p_rooms_idx != -1 and len(r) > p_rooms_idx and r[p_rooms_idx] is not None else '',
                            'size': str(r[p_size_idx]).strip() if p_size_idx != -1 and len(r) > p_size_idx and r[p_size_idx] is not None else ''
                        }

    target_sheets = [upper_sheets['OWNER']] if ('OWNER' in upper_sheets and 'PROPERTY' in upper_sheets) else list(sheet_data.keys())

    cleaned_records = []
    total_original_rows = 0

    for sname in target_sheets:
        rows = sheet_data[sname]
        hdr_idx, hdrs = detect_header_row(rows)
        if hdr_idx == -1 or hdr_idx >= len(rows) - 1:
            continue
        
        data_rows = rows[hdr_idx + 1:]
        total_original_rows += len(data_rows)
        
        lower_hdrs = [re.sub(r'[^a-z0-9]', '', str(h).lower()) for h in hdrs]
        
        # Phone indices
        phone_indices = []
        for i, h in enumerate(lower_hdrs):
            if any(k in h for k in ['mobile', 'phone', 'contact', 'cell', 'whatsapp', 'unnamed2', 'unnamed3', 'number']):
                if not has_neg(h, NEG_PHONE):
                    phone_indices.append(i)
                    
        # Name indices
        name_indices = []
        for i, h in enumerate(lower_hdrs):
            if any(k in h for k in ['owner', 'client', 'customer', 'buyer', 'seller', 'lead', 'name', 'table1']):
                if not has_neg(h, NEG_NAME):
                    name_indices.append(i)
                    
        # Unit indices
        unit_indices = []
        for i, h in enumerate(lower_hdrs):
            if any(k in h for k in ['unit', 'flat', 'villa', 'apartment', 'apt', 'suite', 'door', 'premise', 'townhouse', 'unnamed5']):
                if not has_neg(h, NEG_UNIT):
                    unit_indices.append(i)
        for i, h in enumerate(lower_hdrs):
            if any(k in h for k in ['plotnumber', 'plotno', 'propertyno', 'buildingno', 'plot']):
                if i not in unit_indices and not has_neg(h, NEG_UNIT):
                    unit_indices.append(i)

        # Ref indices
        ref_indices = []
        for i, h in enumerate(lower_hdrs):
            if any(k in h for k in ['project', 'building', 'tower', 'reference', 'pnumber', 'procedure']):
                ref_indices.append(i)
                
        # Location indices
        loc_indices = []
        for i, h in enumerate(lower_hdrs):
            if any(k in h for k in ['masterproject', 'area', 'masterlocation', 'location', 'community', 'district', 'city']):
                if not has_neg(h, NEG_LOC):
                    loc_indices.append(i)
                    
        # Type indices
        type_indices = []
        for i, h in enumerate(lower_hdrs):
            if any(k in h for k in ['propertytype', 'unittype', 'usage', 'type']):
                type_indices.append(i)
                
        # Bed indices
        bed_indices = []
        for i, h in enumerate(lower_hdrs):
            if any(k in h for k in ['rooms', 'bedroom', 'bedrooms', 'beds', 'bed']):
                bed_indices.append(i)
                
        # Size indices
        size_indices = []
        for i, h in enumerate(lower_hdrs):
            if any(k in h for k in ['actualarea', 'totalarea', 'size', 'sqft', 'builtup']):
                size_indices.append(i)

        p_num_idx = next((i for i, h in enumerate(lower_hdrs) if 'pnumber' in h), -1)

        for row in data_rows:
            # 1. Phone extraction with candidate trial
            phone_norm = ""
            raw_phone = None
            for pi in phone_indices:
                if len(row) > pi and row[pi] is not None and str(row[pi]).strip():
                    p_val = str(row[pi]).strip()
                    cand = extract_phone_e164(p_val)
                    if cand:
                        phone_norm = cand
                        raw_phone = p_val
                        break
                    
            # 2. Name
            raw_name = None
            for ni in name_indices:
                if len(row) > ni and row[ni] is not None and str(row[ni]).strip():
                    raw_name = row[ni]
                    break
                    
            # 3. Unit
            raw_unit = None
            for ui in unit_indices:
                if len(row) > ui and row[ui] is not None and str(row[ui]).strip():
                    raw_unit = row[ui]
                    break
                    
            # 4. Ref / Project
            raw_ref = None
            for ri in ref_indices:
                if len(row) > ri and row[ri] is not None and str(row[ri]).strip():
                    raw_ref = row[ri]
                    break
                    
            # 5. Location
            raw_loc = None
            for li in loc_indices:
                if len(row) > li and row[li] is not None and str(row[li]).strip():
                    raw_loc = row[li]
                    break
                    
            # 6. Type
            raw_type = None
            for ti in type_indices:
                if len(row) > ti and row[ti] is not None and str(row[ti]).strip():
                    raw_type = row[ti]
                    break
                    
            # 7. Bedrooms
            raw_beds = None
            for bi in bed_indices:
                if len(row) > bi and row[bi] is not None and str(row[bi]).strip():
                    raw_beds = row[bi]
                    break
                    
            # 8. Size
            raw_size = None
            for si in size_indices:
                if len(row) > si and row[si] is not None and str(row[si]).strip():
                    raw_size = row[si]
                    break

            # If prop_lookup has data for P-NUMBER
            if prop_lookup and p_num_idx != -1 and len(row) > p_num_idx and row[p_num_idx]:
                pkey = str(row[p_num_idx]).strip()
                if pkey in prop_lookup:
                    pinfo = prop_lookup[pkey]
                    if not raw_unit and pinfo['unit']: raw_unit = pinfo['unit']
                    if not raw_ref and pinfo['building']: raw_ref = pinfo['building']
                    if not raw_loc and pinfo['location']: raw_loc = pinfo['location']
                    if not raw_type and pinfo['property_type']: raw_type = pinfo['property_type']
                    if not raw_beds and pinfo['bedrooms']: raw_beds = pinfo['bedrooms']
                    if not raw_size and pinfo['size']: raw_size = pinfo['size']

            # Cell shift detection & healing
            # A. Phone in unit column
            unit_str = str(raw_unit).strip() if raw_unit is not None else ""
            if unit_str:
                d_unit = re.sub(r'\D', '', unit_str)
                if len(d_unit) >= 9 and (unit_str.startswith('05') or unit_str.startswith('+971') or unit_str.startswith('971')):
                    if not phone_norm:
                        phone_cand = extract_phone_e164(unit_str)
                        if phone_cand:
                            phone_norm = phone_cand
                    raw_unit = ""
            # B. Date in unit column
            if unit_str and (re.match(r'^\d{4}[-/]\d{1,2}[-/]\d{1,2}', unit_str) or re.match(r'^\d{1,2}[-/]\d{1,2}[-/]\d{4}', unit_str)):
                raw_unit = ""
                
            # C. Phone in name column
            name_str = str(raw_name).strip() if raw_name is not None else ""
            if name_str and re.match(r'^\+?[\d\s\-_/()]{7,}$', name_str):
                if not phone_norm:
                    phone_cand = extract_phone_e164(name_str)
                    if phone_cand:
                        phone_norm = phone_cand
                raw_name = "Property Owner"

            # Check phone validity
            if not phone_norm:
                continue

            # Normalize fields
            name_clean = clean_owner_name(raw_name)
            unit_clean = clean_unit_number(raw_unit, str(raw_ref or ''))
            
            # Disambiguate reference and location
            ref_clean = str(raw_ref).strip() if raw_ref and str(raw_ref).strip() not in DUMMY_STRINGS else default_ref
            loc_clean = str(raw_loc).strip() if raw_loc and str(raw_loc).strip() not in DUMMY_STRINGS else default_loc
            
            # Check if ref or loc can be mapped to canonical community
            for k, v in COMMUNITY_MAP.items():
                if k in loc_clean.lower():
                    loc_clean = v
                    break

            type_clean = str(raw_type).strip().title() if raw_type and str(raw_type).strip() not in DUMMY_STRINGS else default_type
            if 'villa' in type_clean.lower(): type_clean = "Villa"
            elif 'flat' in type_clean.lower() or 'apt' in type_clean.lower() or 'apartment' in type_clean.lower(): type_clean = "Apartment"
            elif 'town' in type_clean.lower(): type_clean = "Townhouse"
            elif 'pent' in type_clean.lower(): type_clean = "Penthouse"
            elif 'plot' in type_clean.lower() or 'land' in type_clean.lower(): type_clean = "Plot"
            elif 'comm' in type_clean.lower(): type_clean = "Commercial"
            elif 'hotel' in type_clean.lower(): type_clean = "Hotel Apartment"

            beds_clean = clean_bedrooms(raw_beds)
            size_clean = clean_size_sqft(raw_size)

            cleaned_records.append([
                name_clean,
                phone_norm,
                unit_clean,
                ref_clean,
                loc_clean,
                type_clean,
                beds_clean,
                size_clean,
                "Active"
            ])

    # Write output CSV
    with open(out_csv, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow([
            'full_name',
            'phone_number_normalized',
            'unit_number',
            'property_reference',
            'property_location',
            'property_type',
            'bedrooms',
            'size_sqft',
            'status'
        ])
        writer.writerows(cleaned_records)

    duration = round(time.time() - t0, 3)
    status = 'success' if len(cleaned_records) > 0 else 'quarantined'
    return {
        'filename': filename,
        'clean_base_name': clean_base,
        'status': status,
        'original_rows': total_original_rows,
        'cleaned_rows': len(cleaned_records),
        'duration': duration,
        'output_path': out_csv
    }


def main():
    start_all = time.time()
    print("================================================================================")
    print("      DUBAI REAL ESTATE VENDOR DATA CLEANING PIPELINE — MILESTONE 1             ")
    print("================================================================================")
    print(f"Working Directory: {BASE_DIR}")
    print(f"Sources Registry:  {SOURCES_DATA_JS}")
    print(f"Local Data Cache:  {LOCAL_DIR}")
    print(f"Output Directory:  {OUTPUT_DIR}")
    print(f"Manifest Output:   {MANIFEST_PATH}")
    print("--------------------------------------------------------------------------------")

    # Step 1: Parse registered filenames from sourcesData.js
    if not os.path.exists(SOURCES_DATA_JS):
        print(f"ERROR: sourcesData.js not found at {SOURCES_DATA_JS}")
        sys.exit(1)

    with open(SOURCES_DATA_JS, 'r', encoding='utf-8') as f:
        content = f.read()

    raw_matches = re.findall(r'"filename":\s*"([^"]+)"', content)
    print(f"Parsed {len(raw_matches)} catalog records from sourcesData.js.")
    
    # Decode unicode escapes if any
    all_filenames = [json.loads(f'"{fn}"') if '\\u' in fn else fn for fn in raw_matches]
    unique_filenames = sorted(set(all_filenames))
    print(f"Discovered {len(unique_filenames)} unique physical filenames across catalog.")

    # Step 2: Index local filesystem
    local_files_map = {}
    if os.path.exists(LOCAL_DIR):
        for fn in os.listdir(LOCAL_DIR):
            local_files_map[fn] = os.path.join(LOCAL_DIR, fn)
            # also map without U+FFFD replacement char
            if '\ufffd' in fn:
                local_files_map[fn.replace('\ufffd', '-')] = os.path.join(LOCAL_DIR, fn)

    if os.path.exists(CACHE_DIR):
        for fn in os.listdir(CACHE_DIR):
            local_files_map[fn] = os.path.join(CACHE_DIR, fn)

    # Step 3: Check missing files and fetch on-demand
    tasks = []
    missing_count = 0
    downloaded_count = 0

    for fn in unique_filenames:
        target_path = local_files_map.get(fn)
        if not target_path:
            # Try case-insensitive lookup
            fn_low = fn.lower()
            for k, p in local_files_map.items():
                if k.lower() == fn_low and k != 'TownSquare.xlsx': # avoid TownSquare collision with TOWNSQUARE
                    target_path = p
                    break
        
        if not target_path or not os.path.exists(target_path):
            missing_count += 1
            # Download on demand from Supabase storage
            cache_path = os.path.join(CACHE_DIR, fn)
            encoded_fn = requests.utils.quote(fn)
            url = f"{SUPABASE_SOURCES_URL}/{encoded_fn}"
            try:
                print(f"Fetching missing source file: {fn} from Supabase Storage...")
                r = requests.get(url, timeout=30)
                if r.status_code == 200:
                    with open(cache_path, 'wb') as cf:
                        cf.write(r.content)
                    target_path = cache_path
                    downloaded_count += 1
                    local_files_map[fn] = target_path
                    print(f"  Successfully cached {fn} ({len(r.content)} bytes).")
                else:
                    print(f"  Notice: Remote fetch status {r.status_code} for {fn}.")
            except Exception as dl_err:
                print(f"  Remote fetch exception for {fn}: {dl_err}")

        tasks.append((fn, target_path))

    print(f"File Resolution: {len(tasks) - missing_count} pre-cached, {downloaded_count} downloaded on-demand.")

    # Step 4: Execute multi-threaded cleaning pipeline
    print(f"\nLaunching ThreadPoolExecutor with 8 concurrent workers across {len(tasks)} files...")
    manifest = {}
    total_orig = 0
    total_clean = 0
    success_count = 0
    quarantine_count = 0
    error_count = 0

    def worker_job(item):
        fname, fpath = item
        if not fpath or not os.path.exists(fpath):
            clean_base = re.sub(r'\.(xlsx|xls)$', '', fname, flags=re.I)
            out_csv = os.path.join(OUTPUT_DIR, f"{clean_base}_cleaned.csv")
            with open(out_csv, 'w', newline='', encoding='utf-8') as f:
                writer = csv.writer(f)
                writer.writerow(['full_name', 'phone_number_normalized', 'unit_number', 'property_reference', 'property_location', 'property_type', 'bedrooms', 'size_sqft', 'status'])
            return {
                'filename': fname,
                'clean_base_name': clean_base,
                'status': 'quarantined',
                'original_rows': 0,
                'cleaned_rows': 0,
                'duration': 0.0,
                'output_path': out_csv,
                'reason': 'File not found in local cache or storage'
            }
        return process_single_workbook(fpath, fname)

    t_pipeline_start = time.time()
    with ThreadPoolExecutor(max_workers=8) as executor:
        future_to_fn = {executor.submit(worker_job, t): t[0] for t in tasks}
        completed = 0
        total_tasks = len(tasks)
        
        for future in as_completed(future_to_fn):
            fn = future_to_fn[future]
            completed += 1
            try:
                res = future.result()
                manifest[fn] = res
                total_orig += res['original_rows']
                total_clean += res['cleaned_rows']
                
                if res['status'] == 'success':
                    success_count += 1
                elif res['status'] == 'quarantined':
                    quarantine_count += 1
                else:
                    error_count += 1
                    
                if completed % 100 == 0 or completed == total_tasks:
                    pct = (completed / total_tasks) * 100
                    elapsed = time.time() - t_pipeline_start
                    rate = completed / elapsed if elapsed > 0 else 0
                    print(f"[{completed:4d}/{total_tasks:4d}] {pct:5.1f}% | Success: {success_count} | Quarantined: {quarantine_count} | Cleaned rows: {total_clean:,} | Rate: {rate:.1f} files/sec")
            except Exception as exc:
                error_count += 1
                clean_base = re.sub(r'\.(xlsx|xls)$', '', fn, flags=re.I)
                manifest[fn] = {
                    'filename': fn,
                    'clean_base_name': clean_base,
                    'status': 'error',
                    'original_rows': 0,
                    'cleaned_rows': 0,
                    'duration': 0.0,
                    'output_path': os.path.join(OUTPUT_DIR, f"{clean_base}_cleaned.csv"),
                    'reason': str(exc)
                }

    pipeline_duration = time.time() - t_pipeline_start

    # Step 5: Save cleaning manifest
    manifest_payload = {
        'timestamp': time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime()),
        'total_unique_files': len(unique_filenames),
        'total_catalog_entries': len(raw_matches),
        'success_count': success_count,
        'quarantined_count': quarantine_count,
        'error_count': error_count,
        'total_original_rows': total_orig,
        'total_cleaned_rows': total_clean,
        'total_duration_seconds': round(pipeline_duration, 2),
        'files': manifest
    }

    with open(MANIFEST_PATH, 'w', encoding='utf-8') as mf:
        json.dump(manifest_payload, mf, indent=2, ensure_ascii=False)

    print("\n--------------------------------------------------------------------------------")
    print("PIPELINE EXECUTION SUMMARY:")
    print(f"  Total Unique Files Processed: {len(unique_filenames)}")
    print(f"  Successful Files:             {success_count} ({success_count / len(unique_filenames) * 100:.1f}%)")
    print(f"  Quarantined Files:            {quarantine_count} (lock files / 0 phone leads)")
    print(f"  Errors:                       {error_count}")
    print(f"  Total Original Data Rows:     {total_orig:,}")
    print(f"  Total Cleaned Normalized Rows: {total_clean:,}")
    print(f"  Pipeline Processing Time:     {pipeline_duration:.2f} seconds ({len(unique_filenames) / pipeline_duration:.1f} files/sec)")
    print(f"  Manifest Path:                {MANIFEST_PATH}")
    print("================================================================================")

if __name__ == '__main__':
    main()
